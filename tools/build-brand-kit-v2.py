from pathlib import Path
import json, shutil, colorsys, base64, html, zipfile, hashlib
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor, Color, white
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import ImageReader
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor as PC
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
import pypdfium2 as pdfium

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/Axiomotl Branding Kit v2'
QA=ROOT/'outputs/brand-kit-v2-qa'
for folder in ['01 Guidelines','02 Logo','03 Colours and Fonts','04 Icons','05 Office Templates','06 Email','07 Social','08 Print','09 Messaging','10 Developer']:
    (OUT/folder).mkdir(parents=True,exist_ok=True)
QA.mkdir(parents=True,exist_ok=True)
ASSETS=ROOT/'owner-platform/public/site-assets/combined'
LOGO=OUT/'02 Logo/Axiomotl Master Transparent.png'
shutil.copyfile(ASSETS/'axiomotl-logo-master.png',LOGO)
for weight in [400,700,900]:
    shutil.copyfile(ASSETS/f'lato-{weight}.ttf',OUT/f'03 Colours and Fonts/Lato-{weight}.ttf')
    pdfmetrics.registerFont(TTFont(f'Lato{weight}',str(ASSETS/f'lato-{weight}.ttf')))
shutil.copyfile(ASSETS/'OFL.txt',OUT/'03 Colours and Fonts/Lato Licence.txt')
for name in ['Analyse','Design','Decide','Transition','Sustain']:
    src=ROOT/'owner-platform/design/combined/Design-supplied.png' if name=='Design' else ASSETS/f'stages/{name}.png'
    shutil.copyfile(src,OUT/f'04 Icons/{name}.png')

PALETTE={
 'Purple':('#481e72','Primary brand, headings and filled actions'),
 'Night':('#21132f','Hero, presentation cover and dark surfaces'),
 'Teal':('#43aa8c','Progress accents and Design icon centre'),
 'Teal ink':('#276653','Readable teal text on light backgrounds'),
 'Plum':('#7a2c82','Secondary heading emphasis'),
 'Lilac':('#b794d6','Illustration and dark-surface detail'),
 'Lavender':('#e7d6f2','Soft panels and dark-surface buttons'),
 'Paper':('#f7f7fa','Reading surfaces'),
 'White':('#ffffff','Main content and reversed type'),
 'Charcoal':('#333333','Body text'),
 'Blue':('#608599','Supporting diagrams'),
 'Coral':('#e58f80','Supporting illustration, never error text'),
 'Gold':('#d6b35a','Supporting illustration, never warning text')}
def rgb(h):return tuple(int(h[i:i+2],16) for i in (1,3,5))
def luminance(h):
    v=[n/255 for n in rgb(h)]; v=[n/12.92 if n<=.04045 else ((n+.055)/1.055)**2.4 for n in v]
    return sum(a*b for a,b in zip(v,[.2126,.7152,.0722]))
def contrast(a,b):
    x,y=sorted([luminance(a),luminance(b)]);return (y+.05)/(x+.05)
tokens={}
for name,(h,role) in PALETTE.items():
    r,g,b=rgb(h); hue,l,s=colorsys.rgb_to_hls(r/255,g/255,b/255);k=1-max(r,g,b)/255
    cmyk=[0,0,0,100] if k==1 else [round((1-v/255-k)/(1-k)*100) for v in [r,g,b]]+[round(k*100)]
    tokens[name.lower().replace(' ','-')]={'hex':h,'rgb':[r,g,b],'hsl':[round(hue*360),round(s*100),round(l*100)],'cmykApprox':cmyk,'role':role}
(OUT/'03 Colours and Fonts/colour-tokens.json').write_text(json.dumps(tokens,indent=2))
pairs=[('Purple','White'),('Charcoal','White'),('White','Night'),('White','Purple'),('Night','Lavender'),('Teal ink','White'),('Teal','White'),('White','Teal'),('Night','Teal'),('Lilac','Night')]
contrast_rows=[(a,b,contrast(PALETTE[a][0],PALETTE[b][0])) for a,b in pairs]

SECTIONS=[
('Brand foundations',[
('Purpose','Axiomotl Advisory helps organisations connect evidence, decisions and everyday operations. The practice focuses on business analysis, governance and transition in complex environments.'),
('Audience','Organisational leaders, research operations teams, programme sponsors and delivery teams facing fragmented evidence, unclear ownership, system decisions or difficult operational handovers.'),
('Positioning','A small, senior-led advisory practice that turns complex situations into clear requirements, accountable decisions and usable next steps.'),
('Working principles','Understand the actual work. Make ownership explicit. Keep evidence traceable. Leave material the receiving team can use. These are proposed brand principles derived from the current service offer, rather than a claim about a documented founding history.'),
('Name and signature','Write Axiomotl Advisory on first mention and Axiomotl thereafter. Preserve the spelling Axiomotl. The current website headline is “Turn complexity into clarity.” Supporting line: “People. Processes. Progress.” Use one, rather than stacking several slogans.')]),
('The logo',[
('Master artwork','Use the supplied stacked axolotl and AXIOMOTL ADVISORY lockup in 02 Logo. Preserve the original purple artwork, proportions and transparent background. The logo is an image, not a word to be retyped in Lato.'),
('Clear space','Proposed production rule: leave at least one quarter of the symbol width clear on every side. Treat that area as part of the logo when aligning nearby text, photographs or trim edges.'),
('Minimum size','For this detailed stacked mark, use at least 120px overall image width on screen, or 25mm in print after proofing. The 500px-wide raster master reaches about 42mm overall width at 300ppi. Do not upscale it for signage.'),
('Treatments','Use the full-colour transparent master on white or Paper. On dark artwork, place it on a white holding panel with clear space. A true reversed mark, horizontal lockup and simplified small-size symbol are not supplied; commission those from the original artwork before use.'),
('Source discipline','The older geometric SVG and older indigo/teal logo concepts are not part of this kit. No editable vector master is available for the current mark. A PNG embedded in an SVG would still be raster, so this kit does not label one as a vector logo.')]),
('Colour roles',[
('Primary palette','Purple anchors headings and actions; Night establishes depth; Teal signals progress. White and Paper should carry most reading content. Lavender supports soft panels and reversed contexts.'),
('Suggested balance','For documents and information-heavy pages, aim for roughly 70% light neutral, 20% purple/night and 10% accents. A cover may be predominantly Night. This is a composition guide, not a mathematical requirement.'),
('Secondary palette','Plum, Lilac and Blue support hierarchy and diagrams. Coral and Gold are occasional accents. Do not assign a different bright colour to every paragraph, card or heading.'),
('Print specifications','RGB and hex values are the digital reference. CMYK values in the data file are mathematical approximations, not press profiles. Ask the printer for an ICC conversion and physical proof. No Pantone equivalence has been selected.')]),
('Accessible colour',[
('Text pairings','Use Purple or Charcoal on White for standard text. White on Purple or Night also provides strong contrast. Teal ink is the accessible text alternative to the brighter Teal accent.'),
('Accent restrictions','Do not use White text on Teal for small labels, or Teal text on White for body copy. A bright accent is not automatically a readable text colour. Keep a dark label on a teal button.'),
('Interaction states','Proposed UI semantics: success #276653 on #EAF5EF; warning #7A5400 on #FFF4D6; error #A12B36 on #FCEDEF; information #315C78 on #EDF4F8. Pair every state colour with a clear label or icon.'),
('Thresholds','Aim for at least 4.5:1 contrast for normal text and 3:1 for large text and essential interface graphics. Use visible focus rings and avoid encoding meaning through colour alone. Ratios below are calculated from the kit values.')]),
('Typography',[
('Primary family','Use Lato Regular 400, Bold 700 and Black 900. The files and SIL Open Font License are included. Lato is used for both display and body roles on the current site; hierarchy comes from size, weight and spacing.'),
('Digital scale','Hero 48–76px, weight 900, line-height 1.05. Section heading 30–44px, weight 700, line-height 1.15. Subheading 22–26px, weight 700. Body 16–18px, weight 400, line-height 1.6. Small supporting text 13–14px.'),
('Office scale','Use 28pt document titles, 18pt main headings, 13pt subheadings and 11pt body with roughly 1.15–1.3 line spacing. Slides use 30–40pt titles, 18–24pt body and 12–14pt supporting text. Install Lato before editing; Arial is the fallback.'),
('Alignment','Left alignment is the default for reading. The website’s Meet Dr Ramzi paragraphs use justification at the owner’s request. Keep that as a local treatment with hyphenation, rather than justifying narrow mobile text or every business document.'),
('Type care','Use sentence case for headings. Reserve short uppercase labels for navigation or small identifiers. Do not stretch letters, simulate bold, use long all-capital paragraphs or embed essential wording in images.')]),
('Layout and motion',[
('Spacing','Use a 4px base rhythm. Common intervals are 8, 12, 16, 24, 32, 48 and 80px. Current homepage section spacing is 80px desktop and 48px mobile. The requested mobile header-to-hero content gap is 50px; the gap below the workflow cards is 30px.'),
('Page composition','Lead with one clear message, then the evidence or action that supports it. Use generous outer margins, aligned text edges and restrained dividers. Avoid filling every spare area with a badge or decoration.'),
('Digital components','Buttons use clear verbs, comfortable padding and a minimum 44px touch target. Keep body copy in a readable measure of about 45–75 characters. Use a visible keyboard focus state. Decorative effects must not hide text or controls.'),
('Motion','The hero uses restrained depth between existing artwork and background planes. Do not add scroll trapping or empty pinned sections. Keep a complete static composition on mobile and when motion is disabled. Office documents and email remain static.')]),
('Imagery and icons',[
('Photography','Use authentic environments, work materials and human collaboration with a clear relationship to the service. Obtain permission for identifiable people and client material. Do not imply that a stock photograph is a real Axiomotl project.'),
('Illustration','The website uses scientific and system imagery to connect evidence, decisions and outcomes. Keep diagrams structurally meaningful: every connector should express a relationship, sequence or responsibility.'),
('Workflow set','The included website icons are Analyse, Design, Decide, Transition and Sustain. The Design icon is the user-supplied transparent teal-centre version. Keep them on consistent visual bounds; never distort their proportions.'),
('Where to use them','Use the dimensional PNG icons in marketing and presentations at small sizes. In dense documents, use editable labels, tables or simple line diagrams instead of tiny shaded pictures. Keep the original icon set together.'),
('Rights and handling','This kit reuses client-supplied logo and icon artwork. It does not independently establish ownership or a licence for those assets. Lato has its included OFL. No third-party stock photography is redistributed.')]),
('Voice and tone',[
('Voice','Clear: name the decision and next step. Practical: show what the work produces. Informed: connect a recommendation to evidence. Calm: be direct about constraints without exaggeration.'),
('Vocabulary','Prefer evidence, requirements, ownership, decision, governance, readiness and handover. Explain business-as-usual before using BAU. Avoid unqualified claims such as world-leading, guaranteed transformation and best-in-class.'),
('Style','Use Australian English: organisation, prioritise, programme where appropriate. Address the reader as you. Keep headings in sentence case. Use short paragraphs and specific verbs. Preserve the formal names of services.'),
('Tone by situation','Marketing should be inviting and concrete. Proposals should define scope and responsibilities. Findings should distinguish evidence from interpretation. Support messages should explain what happened and what to do next.'),
('Before and after','Instead of “Unlock unparalleled enterprise transformation”, write “Agree the requirements and decisions your delivery team needs.” Instead of “An error occurred”, write “Your changes were not saved. Check your connection and try again.”')]),
('Messaging library',[
('One sentence','Axiomotl Advisory provides senior business analysis and governance support to help organisations turn complex work into clear decisions and practical next steps.'),
('Short description','Axiomotl Advisory connects evidence, requirements and decision ownership across complex organisational change. Led by Dr Ramzi Abbassi, the practice supports diagnostics, decision architecture, operational transition and embedded senior business analysis.'),
('Services','Business & Governance Diagnostic; Requirements & Decision Architecture; BAU Transition & Operating Model; Embedded Principal BA Advisory.'),
('Practitioner biography','Dr Ramzi Abbassi, BMedSc (Hons), PhD, CCBA, brings experience spanning scientific research, laboratory digitalisation and enterprise business analysis. His published work discusses governance, operating models, data migration, rollout and transition to everyday operations.'),
('Claims boundary','Public biographies and named organisations describe practitioner background, not Axiomotl client endorsements. Do not invent client logos, performance statistics, founding dates, awards, qualifications or current employer claims.')]),
('Business applications',[
('Word templates','The kit includes letterhead, an engagement proposal, a decision brief and meeting notes. Replace every bracketed prompt before sharing. Keep scope, fees and acceptance details specific to the engagement; the proposal is not a legal contract.'),
('Presentation template','Eight editable widescreen slides cover an introduction, context, services, method, findings, decision, action plan and close. Replace sample prompts and duplicate slides as needed. Text, diagrams and shapes remain editable.'),
('Email','Use the light and dark-accent HTML signatures as starting points, plus the plain-text version. Do not include passwords, tracking pixels or an unconfirmed phone number. The signature uses a text wordmark to avoid blocked external images.'),
('Social','Editable SVG templates are supplied with PNG exports for a square post, portrait post, landscape update and profile banner. Text remains live in SVG. Check platform crop previews before posting; dimensions are production formats, not guarantees of platform safe areas.'),
('Print','The two-sided business card PDF uses 90 × 55mm trim with 3mm bleed. Letterhead is A4. Proof print colour and logo detail with the printer. The business card is a design proof and does not carry a invented address, phone number or ABN.')]),
('Usage boundaries',[
('Logo','Do preserve the supplied master, clear space and proportions. Do not recolour it, stretch it, rotate it or add new gradients. Put it on a white holding panel if the background would reduce legibility.'),
('Colour','Do use strong text pairings and limit accents. Do not use white small text on Teal or rely on colour alone to communicate a status.'),
('Typography','Do use real Lato weights and a clear hierarchy. Do not use tiny text to force content into a space. Shorten the copy or add a page.'),
('Composition','Do align elements and group related material. Do not stretch photos or let decorative lines run through labels. Keep social text inside the template guides.'),
('Voice','Do state the practical output and the evidence. Do not imply guaranteed results, medical advice, client endorsement or certification that has not been verified.')]),
('Handoff and maintenance',[
('Version','Version 2 • 21 September 2026. This kit records the current purple-and-teal website identity and supersedes the earlier indigo/teal working kit for new material. Existing live website content has not been changed by creating this package.'),
('Authority','Existing supplied logo, website colours, Lato typography and the supplied Design icon are retained. Clear-space, minimum-size, semantic-colour and application rules are proposed production standards for owner review.'),
('Outstanding master artwork','Ask the original designer for editable vector artwork and approved reversed, horizontal and small-format logo variants before ordering large-format printing, embroidery or signage. No substitute logo is represented as approved in this kit.'),
('Release checklist','Replace all bracketed prompts. Check contact details and links. Confirm claim accuracy. Review mobile crops. Proof print colour. Confirm accessibility. Save a dated copy and note which files changed.'),
('Sources','Identity: current website and client-supplied logo/icon files. Professional background: SciSure interview dated 6 February 2024 and Ramzi’s implementation article dated 10 August 2026. Source URLs and asset hashes are included in the kit. These references support the biography, not a commercial endorsement.')])]

# Canonical text and offline HTML guide.
md=['# Axiomotl Advisory Brand Guidelines','Version 2 | 21 September 2026','Current identity with proposed production rules for owner review.']
for title,blocks in SECTIONS:
    md += ['\n## '+title]+['\n### '+h+'\n'+p for h,p in blocks]
md+=['\n## Palette']+[f'- {n}: {h} — {r}' for n,(h,r) in PALETTE.items()]
md+=['\n## Contrast']+[f'- {a} on {b}: {r:.2f}:1' for a,b,r in contrast_rows]
(OUT/'01 Guidelines/Brand Guidelines.md').write_text('\n'.join(md),encoding='utf-8')
html_sections=''.join('<section id="s'+str(i)+'"><h2>'+html.escape(t)+'</h2>'+''.join('<h3>'+html.escape(h)+'</h3><p>'+html.escape(p)+'</p>' for h,p in blocks)+'</section>' for i,(t,blocks) in enumerate(SECTIONS))
swatches=''.join(f'<div class="swatch"><i style="background:{h}"></i><b>{n}</b><span>{h}</span><p>{r}</p></div>' for n,(h,r) in PALETTE.items())
nav=''.join(f'<a href="#s{i}">{t}</a>' for i,(t,b) in enumerate(SECTIONS))
web=f'''<!doctype html><html lang="en-AU"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Axiomotl brand guidelines</title><style>@font-face{{font-family:Lato;src:url('../03 Colours and Fonts/Lato-400.ttf')}}@font-face{{font-family:Lato;src:url('../03 Colours and Fonts/Lato-700.ttf');font-weight:700}}*{{box-sizing:border-box}}body{{margin:0;font-family:Lato,Arial,sans-serif;color:#333;background:#f7f7fa;line-height:1.65}}header{{padding:60px max(24px,calc((100% - 1100px)/2));background:#21132f;color:white}}header img{{width:130px;background:white;border-radius:8px}}h1{{font-size:clamp(36px,6vw,64px);line-height:1.08;max-width:800px}}main{{max-width:1100px;margin:auto;padding:32px 24px}}nav{{display:flex;gap:12px 22px;flex-wrap:wrap;border-bottom:1px solid #c8b8d3;padding-bottom:28px}}a{{color:#481e72}}section{{padding:35px 0;border-bottom:1px solid #c8b8d3}}h2{{color:#481e72;font-size:32px}}h3{{font-size:19px;margin-bottom:6px}}p{{max-width:78ch;margin-top:0}}.palette{{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:20px}}.swatch i{{display:block;height:90px;border:1px solid #ddd}}.swatch b,.swatch span{{display:block}}.swatch p{{font-size:13px}}.icons{{display:flex;flex-wrap:wrap;gap:30px}}.icons img{{width:100px;height:100px;object-fit:contain}}@media print{{nav{{display:none}}section{{break-inside:avoid}}}}</style><header><img src="../02 Logo/Axiomotl Master Transparent.png" alt="Axiomotl Advisory"><h1>Axiomotl Advisory<br>Brand guidelines</h1><p>Version 2 · 21 September 2026</p></header><main><nav>{nav}</nav>{html_sections}<section><h2>Colour reference</h2><div class="palette">{swatches}</div></section><section><h2>Workflow icons</h2><div class="icons">{''.join(f'<figure><img src="../04 Icons/{n}.png" alt=""><figcaption>{n}</figcaption></figure>' for n in ['Analyse','Design','Decide','Transition','Sustain'])}</div></section></main></html>'''
(OUT/'01 Guidelines/Brand Guidelines.html').write_text(web,encoding='utf-8')

# PDF guide, carefully fixed-page layout.
W,H=595.28,841.89
c=canvas.Canvas(str(OUT/'01 Guidelines/Axiomotl Brand Guidelines.pdf'),pagesize=(W,H))
c.setTitle('Axiomotl Advisory Brand Guidelines v2')
def para(text,x,y,width,size=11,color='#333333',leading=None):
    st=ParagraphStyle('p',fontName='Lato400',fontSize=size,leading=leading or size*1.42,textColor=HexColor(color))
    p=Paragraph(html.escape(text),st);_,hh=p.wrap(width,H);p.drawOn(c,x,y-hh);return y-hh
def page_base(title,num):
    c.setFillColor(HexColor('#ffffff'));c.rect(0,0,W,H,fill=1,stroke=0)
    c.setFillColor(HexColor('#481e72'));c.rect(0,H-12,W,12,fill=1,stroke=0)
    c.setFont('Lato700',10);c.drawString(44,H-44,'AXIOMOTL ADVISORY / BRAND GUIDELINES')
    c.setFont('Lato900',30);c.drawString(44,H-95,title)
    c.setStrokeColor(HexColor('#e7d6f2'));c.line(44,48,W-44,48)
    c.setFont('Lato400',9);c.setFillColor(HexColor('#56515c'));c.drawString(44,31,'Version 2 · September 2026');c.drawRightString(W-44,31,str(num))
def finish():c.showPage()
c.setFillColor(HexColor('#21132f'));c.rect(0,0,W,H,fill=1,stroke=0)
c.setFillColor(white);c.roundRect(44,H-207,145,145,8,fill=1,stroke=0);c.drawImage(str(LOGO),50,H-201,133,133,mask='auto')
c.setFillColor(white);c.setFont('Lato900',47);c.drawString(44,460,'Axiomotl');c.drawString(44,403,'Advisory')
c.setFont('Lato400',26);c.drawString(44,341,'Brand guidelines and toolkit')
c.setFillColor(HexColor('#43aa8c'));c.rect(44,286,86,6,fill=1,stroke=0)
para('The identity, language and practical tools for clear, consistent communication.',44,248,450,17,'#e7d6f2')
para('Version 2 · 21 September 2026\nCurrent identity with proposed production standards for owner review.',44,120,450,11,'#e7d6f2');finish()
for idx,(title,blocks) in enumerate(SECTIONS,2):
    page_base(title,idx);y=H-131
    for heading,text in blocks:
        c.setFillColor(HexColor('#481e72'));c.setFont('Lato700',13);c.drawString(44,y,heading);y-=10
        y=para(text,44,y,W-88,10.8);y-=21
    if y<65:raise RuntimeError(f'Guide overflow {title}: {y}')
    finish()
page_base('Palette reference',14)
for i,(name,(h,role)) in enumerate(PALETTE.items()):
    row,col=divmod(i,3);x=44+col*173;y=H-137-row*133
    c.setFillColor(HexColor(h));c.rect(x,y-48,150,48,fill=1,stroke=0)
    c.setFillColor(HexColor('#333333'));c.setFont('Lato700',11);c.drawString(x,y-65,name+' '+h)
    para('RGB '+', '.join(map(str,rgb(h))),x,y-76,154,9)
    para(role,x,y-92,154,8.5)
finish()
page_base('Contrast reference',15);y=H-142
for a,b,ratio in contrast_rows:
    c.setFillColor(HexColor(PALETTE[b][0]));c.rect(44,y-29,245,34,fill=1,stroke=0)
    c.setFillColor(HexColor(PALETTE[a][0]));c.setFont('Lato700',12);c.drawString(53,y-15,f'{a} on {b}')
    c.setFillColor(HexColor('#333333'));c.setFont('Lato400',11);c.drawString(310,y-15,f'{ratio:.2f}:1  '+('Normal text passes' if ratio>=4.5 else 'Avoid normal text'));y-=47
para('Calculated WCAG contrast ratios. The document specifies preferred pairings; it is not a certification of every future application.',44,175,W-88,11)
finish()
page_base('The identity in use',16)
c.setFillColor(HexColor('#f7f7fa'));c.rect(44,489,230,210,fill=1,stroke=0);c.drawImage(str(LOGO),81,510,155,155,mask='auto')
para('Master on a light surface. Preserve transparent edges and the complete lockup.',44,471,230,10)
c.setFillColor(HexColor('#21132f'));c.rect(299,489,252,210,fill=1,stroke=0)
c.setFillColor(white);c.roundRect(345,515,155,155,8,fill=1,stroke=0);c.drawImage(str(LOGO),351,521,143,143,mask='auto')
para('A white holding panel keeps the supplied purple master legible on Night.',299,471,250,10)
for i,n in enumerate(['Analyse','Design','Decide','Transition','Sustain']):
    c.drawImage(str(OUT/f'04 Icons/{n}.png'),48+i*101,270,80,80,preserveAspectRatio=True,anchor='c',mask='auto')
    c.setFillColor(HexColor('#481e72'));c.setFont('Lato700',10);c.drawCentredString(88+i*101,250,n)
para('Use the five workflow icons as a coherent set. The current Design artwork has a teal centre and transparent background.',44,212,W-88,11)
para('The next files put the system to work: editable Word and PowerPoint templates, SVG social layouts with PNG exports, email signatures, and business card artwork.',44,140,W-88,11)
finish();c.save()

# Word templates with ordinary, editable structure.
def doc_base(title,intro):
    d=Document();sec=d.sections[0];sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929);sec.top_margin=Inches(.8);sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
    for nm in ['Normal','Title','Heading 1','Heading 2']:
        st=d.styles[nm];st.font.name='Lato';st.font.size=Pt(11 if nm=='Normal' else 28 if nm=='Title' else 18 if nm=='Heading 1' else 13)
        st.font.color.rgb=RGBColor.from_string('000000' if nm=='Title' else '481E72' if nm.startswith('Heading') else '333333')
    d.styles['Normal'].paragraph_format.space_after=Pt(8)
    h=sec.header.paragraphs[0];h.add_run().add_picture(str(LOGO),width=Inches(.82));h.add_run('   AXIOMOTL ADVISORY').bold=True
    f=sec.footer.paragraphs[0];f.add_run('Axiomotl Advisory  |  hello@axiomotl.com.au').font.size=Pt(9)
    d.add_paragraph(title,'Title');d.add_paragraph(intro);return d
def table(d,heads,rows):
    t=d.add_table(rows=1,cols=len(heads));t.style='Light Shading Accent 1'
    for cell,h in zip(t.rows[0].cells,heads):cell.text=h
    for row in rows:
        for cell,txt in zip(t.add_row().cells,row):cell.text=txt
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:r.font.size=Pt(10)
    return t
docs=[]
d=doc_base('Letterhead','[Date]\n[Recipient name]\n[Organisation]\n[Address]');d.add_paragraph('Subject','Heading 1');d.add_paragraph('[State the reason for writing and the action or decision needed.]');d.add_paragraph('Dear [Name]');d.add_paragraph('[Provide the relevant context in short paragraphs. Explain the next step and who owns it.]');d.add_paragraph('Kind regards\nDr Ramzi Abbassi\nAxiomotl Advisory');docs.append(('Letterhead',d))
d=doc_base('Engagement proposal','Prepared for [Client organisation] by Axiomotl Advisory. This proposal defines the work to be agreed, the intended outputs and the decisions required before starting.');d.add_paragraph('Context and objective','Heading 1');d.add_paragraph('[Describe the situation, the decision ahead and the desired operational outcome.]');d.add_paragraph('Scope and outputs','Heading 1');table(d,['Work included','Output'],[['[Activity]','[Deliverable and acceptance criteria]'],['[Activity]','[Deliverable and acceptance criteria]']]);d.add_paragraph('Working arrangements','Heading 1');d.add_paragraph('[List the sponsor, key stakeholders, access requirements, assumptions and exclusions.]');d.add_paragraph('Timing and commercial details','Heading 1');d.add_paragraph('[Insert milestones, fees, GST treatment, payment arrangements and the applicable engagement agreement. Confirm all figures before sharing.]');d.add_paragraph('Next step','Heading 1');d.add_paragraph('[Name the person responsible for confirming scope and the date for a decision.]');docs.append(('Engagement Proposal',d))
d=doc_base('Decision brief','Prepared for [Decision owner] on [Date]. Use this brief to connect the decision to its evidence, options and next steps.');d.add_paragraph('Decision required','Heading 1');d.add_paragraph('[State the decision and when it is needed.]');d.add_paragraph('Evidence and constraints','Heading 1');d.add_paragraph('[Separate verified facts from assumptions. Link to the source material.]');d.add_paragraph('Options','Heading 1');table(d,['Option','Benefits and constraints'],[['[Option A]','[Benefits, risks, dependencies]'],['[Option B]','[Benefits, risks, dependencies]']]);d.add_paragraph('Recommendation','Heading 1');d.add_paragraph('[Explain which option is recommended, why and what would change the recommendation.]');d.add_paragraph('Ownership and follow through','Heading 1');table(d,['Action','Owner','Due'],[['[Next step]','[Name]','[Date]'],['[Review point]','[Name]','[Date]']]);docs.append(('Decision Brief',d))
d=doc_base('Meeting notes','[Meeting name]  |  [Date and time]\nParticipants [Names]  |  Facilitator [Name]');d.add_paragraph('Purpose','Heading 1');d.add_paragraph('[Describe what the meeting needs to resolve.]');d.add_paragraph('Discussion and evidence','Heading 1');d.add_paragraph('[Record the material points, evidence and unresolved questions.]');d.add_paragraph('Decisions','Heading 1');table(d,['Decision','Owner','Reason'],[['[Decision]','[Name]','[Evidence or rationale]']]);d.add_paragraph('Actions','Heading 1');table(d,['Action','Owner','Due'],[['[Action]','[Name]','[Date]'],['[Action]','[Name]','[Date]']]);d.add_paragraph('Next meeting','Heading 1');d.add_paragraph('[Date, purpose and preparation needed.]');docs.append(('Meeting Notes',d))
for name,d in docs:
    path=OUT/f'05 Office Templates/{name}.docx';d.save(path)
    # Update the Office theme so editable tables use brand colours.
    with zipfile.ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
    import re
    theme=parts['word/theme/theme1.xml'].decode()
    for key,val in [('accent1','481E72'),('accent2','43AA8C'),('accent3','7A2C82'),('accent4','608599'),('accent5','B794D6'),('accent6','D6B35A')]:
        theme=re.sub(r'(<a:'+key+r'>).*?(</a:'+key+r'>)',r'\1<a:srgbClr val="'+val+r'"/>\2',theme)
    parts['word/theme/theme1.xml']=theme.encode()
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for n,b in parts.items():z.writestr(n,b)

# Editable eight-slide presentation.
prs=Presentation();prs.slide_width=PI(13.333);prs.slide_height=PI(7.5)
def textbox(sl,text,x,y,w,h,size=22,color='333333',bold=False):
    box=sl.shapes.add_textbox(PI(x),PI(y),PI(w),PI(h));tf=box.text_frame;tf.word_wrap=True
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Lato';p.font.size=PP(size);p.font.bold=bold;p.font.color.rgb=PC.from_string(color)
    return box
def slide(title,dark=False):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=PC.from_string('21132F' if dark else 'FFFFFF')
    textbox(s,title,.65,.65,10.8,1.15,34,'FFFFFF' if dark else '481E72',True)
    textbox(s,'Axiomotl Advisory',.7,7.0,8,.25,11,'E7D6F2' if dark else '56515C')
    return s
s=slide('[Presentation title]',True);textbox(s,'[Client or project]\n[Presenter] · [Date]',.7,2.35,8,1.4,24,'E7D6F2');shape=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,PI(10.25),PI(4.55),PI(2.1),PI(2.1));shape.fill.solid();shape.fill.fore_color.rgb=PC(255,255,255);shape.line.fill.background();s.shapes.add_picture(str(LOGO),PI(10.32),PI(4.62),width=PI(1.96))
s=slide('The situation and decision');textbox(s,'What is happening',.7,2,5.5,.5,24,'481E72',True);textbox(s,'[Describe the current situation.]\n[Identify the evidence and constraints.]',.7,2.7,5.5,2.3);textbox(s,'What needs deciding',7.0,2,5.5,.5,24,'481E72',True);textbox(s,'[State the decision required.]\n[Name the owner and timing.]',7,2.7,5.5,2.3)
s=slide('Four ways to work together')
services=['Business & Governance Diagnostic','Requirements & Decision Architecture','BAU Transition & Operating Model','Embedded Principal BA Advisory']
for i,t in enumerate(services):
    x=.7+(i%2)*6.2;y=2+(i//2)*2.1;textbox(s,t,x,y,5.5,.8,25,'481E72',True);textbox(s,'[Describe the relevant scope and output.]',x,y+.9,5.4,.7,18)
s=slide('From analysis to sustained operations')
for i,n in enumerate(['Analyse','Design','Decide','Transition','Sustain']):
    x=.7+i*2.52;s.shapes.add_picture(str(OUT/f'04 Icons/{n}.png'),PI(x+.5),PI(2.15),width=PI(1.0));textbox(s,n,x,3.4,2.3,.5,23,'481E72',True);textbox(s,'[Output or decision]',x,4.1,2.25,.8,17)
s=slide('What the evidence shows')
for i,t in enumerate(['Finding','Evidence','Implication']):textbox(s,t,.7+i*4.2,2,3.8,.5,24,'481E72',True);textbox(s,'[Add a concise, sourced statement.]',.7+i*4.2,2.8,3.8,2.5,22)
s=slide('The decision and its trade offs');textbox(s,'Recommended option',.7,2,5.6,.5,24,'481E72',True);textbox(s,'[State the recommendation.]\n[Explain why the evidence supports it.]',.7,2.8,5.6,2.8);textbox(s,'Risks and dependencies',7,2,5.5,.5,24,'481E72',True);textbox(s,'[Name the constraints.]\n[Explain how they will be managed.]',7,2.8,5.5,2.8)
s=slide('Actions and ownership');tbl=s.shapes.add_table(4,3,PI(.7),PI(2),PI(11.9),PI(3.4)).table
for j,h in enumerate(['Action','Owner','Due']):tbl.cell(0,j).text=h
for i in range(1,4):
    for j,t in enumerate(['[Next step]','[Name]','[Date]']):tbl.cell(i,j).text=t
for i,row in enumerate(tbl.rows):
    for cell in row.cells:
        cell.fill.solid();cell.fill.fore_color.rgb=PC.from_string('481E72' if i==0 else 'F7F7FA')
        for p in cell.text_frame.paragraphs:p.font.name='Lato';p.font.size=PP(20);p.font.color.rgb=PC.from_string('FFFFFF' if i==0 else '333333')
s=slide('Agree the next step',True);textbox(s,'[Decision, owner and date]',.7,2.2,11,1.2,32,'E7D6F2');textbox(s,'Dr Ramzi Abbassi\nhello@axiomotl.com.au',.7,4.3,10,1.3,24,'FFFFFF');prs.save(OUT/'05 Office Templates/Axiomotl Presentation Template.pptx')

# Email signatures and reusable message templates.
for name,bg,fg,accent in [('Light','FFFFFF','333333','481E72'),('Dark accent','21132F','FFFFFF','E7D6F2')]:
    sig=f'''<!doctype html><html><meta charset="utf-8"><title>Axiomotl email signature</title><body><table role="presentation" cellpadding="0" cellspacing="0" style="font-family:Arial,sans-serif;font-size:14px;background:#{bg};color:#{fg};padding:18px;border-left:4px solid #43aa8c"><tr><td style="font-weight:bold;font-size:17px">Dr Ramzi Abbassi</td></tr><tr><td style="padding:4px 0 12px">BMedSc (Hons), PhD, CCBA</td></tr><tr><td style="font-weight:bold;letter-spacing:1px;color:#{accent}">AXIOMOTL ADVISORY</td></tr><tr><td style="padding-top:8px"><a href="mailto:hello@axiomotl.com.au" style="color:#{accent}">hello@axiomotl.com.au</a></td></tr><tr><td style="padding-top:5px"><a href="https://axiomotl-advisory.vercel.app/" style="color:#{accent}">Visit our website</a></td></tr></table></body></html>'''
    (OUT/f'06 Email/Signature {name}.html').write_text(sig,encoding='utf-8')
(OUT/'06 Email/Signature Plain Text.txt').write_text('Dr Ramzi Abbassi\nBMedSc (Hons), PhD, CCBA\nAxiomotl Advisory\nhello@axiomotl.com.au\nhttps://axiomotl-advisory.vercel.app/\n',encoding='utf-8')
(OUT/'06 Email/Email Templates.md').write_text('''# Axiomotl email templates

## Introduction
Subject: A conversation about [specific issue]

Hi [Name],

I understand your team is working through [situation]. Axiomotl can help clarify the requirements, decisions and ownership needed for the next step.

Would a short conversation about [specific question] be useful?

Kind regards,
Ramzi

## Following a meeting
Subject: Next steps from our discussion

Hi [Name],

Thank you for the discussion. We agreed that [decision]. The next step is [action], owned by [person], by [date].

I have attached [document] for review. Please let me know if anything differs from your understanding.

Kind regards,
Ramzi

## Requesting feedback
Subject: [Document] ready for your review

Hi [Name],

The draft [document] is ready. Please focus your review on [specific points], particularly any gaps in the evidence, responsibilities or proposed next steps.

Could you send your comments by [date]?

Kind regards,
Ramzi
''',encoding='utf-8')

# Editable SVG social layouts and exact raster exports from matching PDF primitives.
logo64=base64.b64encode(LOGO.read_bytes()).decode()
formats=[('Square post',1080,1080,['Turn complexity','into clarity.']),('Portrait post',1080,1350,['Clear requirements.','Accountable decisions.']),('Landscape update',1200,630,['Make the next','decision clear.']),('Profile banner',1584,396,['Axiomotl Advisory','Evidence. Decisions. Operations.'])]
for name,w,h,lines in formats:
    pad=round(w*.065);f=round(w*.053 if name!='Profile banner' else 53);top=int(h*.45); accent_y=95 if name=='Profile banner' else h-pad-75
    lx=w-pad-150;ly=pad
    els=[f'<rect width="{w}" height="{h}" fill="#21132f"/>',f'<rect x="{lx}" y="{ly}" width="150" height="150" rx="8" fill="#fff"/>',f'<image href="data:image/png;base64,{logo64}" x="{lx+6}" y="{ly+6}" width="138" height="138"/>']
    for i,line in enumerate(lines):els.append(f'<text x="{pad}" y="{top+i*(f+16)}" font-family="Lato,Arial,sans-serif" font-size="{f}" font-weight="700" fill="{ "#ffffff" if i==0 else "#e7d6f2" }">{html.escape(line)}</text>')
    els += [f'<rect x="{pad}" y="{accent_y}" width="90" height="6" fill="#43aa8c"/>',f'<text x="{pad}" y="{h-pad}" font-family="Lato,Arial,sans-serif" font-size="{max(19,int(w*.022))}" fill="#ffffff">hello@axiomotl.com.au</text>']
    if name=='Profile banner':
        # Keep the left quarter clear of primary content for a typical avatar overlap.
        els=[el.replace(f'x="{pad}"',f'x="{int(w*.29)}"') for el in els]
    svg=f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" width="{w}" height="{h}" viewBox="0 0 {w} {h}">'+''.join(els)+'</svg>'
    (OUT/f'07 Social/{name}.svg').write_text(svg,encoding='utf-8')
    temp=QA/f'{name}.pdf';sc=canvas.Canvas(str(temp),pagesize=(w,h));sc.setFillColor(HexColor('#21132f'));sc.rect(0,0,w,h,fill=1,stroke=0);sc.setFillColor(white);sc.roundRect(lx,h-ly-150,150,150,8,fill=1,stroke=0);sc.drawImage(str(LOGO),lx+6,h-ly-144,138,138,mask='auto')
    x=int(w*.29) if name=='Profile banner' else pad
    for i,line in enumerate(lines):sc.setFont('Lato700',f);sc.setFillColor(HexColor('#ffffff' if i==0 else '#e7d6f2'));sc.drawString(x,h-top-i*(f+16),line)
    sc.setFillColor(HexColor('#43aa8c'));sc.rect(x,h-accent_y-6,90,6,fill=1,stroke=0);sc.setFillColor(white);sc.setFont('Lato400',max(19,int(w*.022)));sc.drawString(x,pad,'hello@axiomotl.com.au');sc.save()
    pdf=pdfium.PdfDocument(str(temp));pdf[0].render(scale=1).to_pil().save(OUT/f'07 Social/{name}.png');pdf.close()

# Two-sided business card with PDF trim and bleed boxes.
mm=72/25.4;bw,bh=96*mm,61*mm
bc=canvas.Canvas(str(OUT/'08 Print/Business Card Proof.pdf'),pagesize=(bw,bh))
for side in [0,1]:
    bc.setTrimBox((3*mm,3*mm,93*mm,58*mm));bc.setBleedBox((0,0,bw,bh))
    bc.setFillColor(HexColor('#ffffff' if side==0 else '#21132f'));bc.rect(0,0,bw,bh,fill=1,stroke=0)
    if side==0:bc.drawImage(str(LOGO),(bw-35*mm)/2,(bh-35*mm)/2,35*mm,35*mm,mask='auto')
    else:
        bc.setFillColor(white);bc.setFont('Lato700',13);bc.drawString(9*mm,43*mm,'Dr Ramzi Abbassi');bc.setFont('Lato400',8);bc.drawString(9*mm,36*mm,'BMedSc (Hons), PhD, CCBA');bc.setFillColor(HexColor('#43aa8c'));bc.rect(9*mm,29*mm,20*mm,1.2*mm,fill=1,stroke=0);bc.setFillColor(white);bc.setFont('Lato700',9);bc.drawString(9*mm,22*mm,'AXIOMOTL ADVISORY');bc.setFont('Lato400',8.5);bc.drawString(9*mm,15*mm,'hello@axiomotl.com.au')
    bc.showPage()
bc.save()
(OUT/'08 Print/Print Production Notes.txt').write_text('Business card: 90 x 55mm trim, 3mm bleed on all sides, 96 x 61mm page. PDF has TrimBox and BleedBox. Two pages represent front and back. RGB design proof, not a certified PDF/X press file. Convert with the printer profile and proof the logo before ordering.\nA4 letterhead is in Office Templates. No phone number, ABN or postal address was provided. Add these only after confirmation.\n',encoding='utf-8')

css=':root {\n'+''.join(f'  --axiomotl-{k}: {v["hex"]};\n' for k,v in tokens.items())+'  --axiomotl-font: Lato, Arial, sans-serif;\n  --axiomotl-radius: 6px;\n  --axiomotl-section: 80px;\n}\n@media (max-width: 700px) { :root { --axiomotl-section: 48px; } }\n'
(OUT/'10 Developer/brand-tokens.css').write_text(css)
shutil.copyfile(OUT/'03 Colours and Fonts/colour-tokens.json',OUT/'10 Developer/brand-tokens.json')
(OUT/'03 Colours and Fonts/Palette.gpl').write_text('GIMP Palette\nName: Axiomotl Advisory v2\nColumns: 4\n#\n'+''.join(f'{r:3} {g:3} {b:3}\t{name}\n' for name,(h,_) in PALETTE.items() for r,g,b in [rgb(h)]))
(OUT/'03 Colours and Fonts/Contrast.csv').write_text('Foreground,Background,Ratio,Normal text\n'+''.join(f'{a},{b},{r:.2f},{"Pass" if r>=4.5 else "Fail"}\n' for a,b,r in contrast_rows))
(OUT/'09 Messaging/Copy Library.md').write_text('\n'.join(['# Axiomotl messaging']+['\n## '+h+'\n'+p for title,blocks in SECTIONS if title in ['Voice and tone','Messaging library'] for h,p in blocks]),encoding='utf-8')
(OUT/'09 Messaging/Source Notes.md').write_text('''# Sources and decisions

Identity reference: https://axiomotl-advisory.vercel.app/
Current practice profile: https://axiomotl-advisory.vercel.app/practice
Professional biography: https://www.scisure.com/blog/implement-scientific-management-platforms-at-scale (10 August 2026)
Professional interview: https://www.scisure.com/blog/client-to-team-member-ramzi-abbassi-phd (6 February 2024)

The sources were reviewed during the website build in this task. Employer appointments are intentionally omitted from reusable biographies because source dates vary. The kit does not add a founding date, awards, numerical outcomes or client endorsements.

Logo: supplied 500 x 500 transparent purple master, as stored in the website workspace. It is not the older blue/indigo logo and not the geometric SVG concept.
Design icon: the user-supplied icon.png copied during the current task. Other icons: the live website workflow set.

Proposed standards: clear space, minimum sizes, semantic colours, office layouts and social layouts. Owner review is appropriate before treating those as mandatory organisation-wide policy.
''',encoding='utf-8')
(OUT/'02 Logo/Logo Usage.txt').write_text('Use Axiomotl Master Transparent.png unchanged. 500 x 500px RGBA. Maintain aspect ratio and transparency. Minimum proposed screen width 120px; minimum print width 25mm after proof. Maximum overall width about 42mm at 300ppi. On dark backgrounds use a white holding panel. No final vector, reversed, horizontal or symbol-only variant was supplied. Do not use the older geometric SVG as this logo.\n',encoding='utf-8')
(OUT/'04 Icons/README.txt').write_text('Current website workflow icons. Preserve aspect ratio. Design.png is the latest client-supplied teal-centre icon. These files have transparent backgrounds. They are raster marketing icons, not an editable vector family.\n',encoding='utf-8')
(OUT/'07 Social/README.txt').write_text('Open SVG files in a vector editor to edit live text and shapes. Install Lato before editing. Logos are embedded raster artwork. PNG files are ready to use. Layout dimensions: square 1080x1080; portrait 1080x1350; landscape 1200x630; profile banner 1584x396. Preview cropping in the target platform before publication. Profile banner keeps the left quarter clear of primary content.\n',encoding='utf-8')
(OUT/'06 Email/README.txt').write_text('Open the HTML signature, select the rendered signature and paste into your email application signature editor. Use the plain-text version where formatting is unsupported. Test a message to yourself in your email client before sending externally. No email has been sent. No passwords or authentication details are included.\n',encoding='utf-8')
(OUT/'05 Office Templates/README.txt').write_text('Install the bundled Lato fonts before editing. DOCX and PPTX files are editable starting documents. Replace bracketed prompts before use. In Office, use Save As to make DOTX or POTX templates if preferred. The proposal is a scoping document, not a substitute for the applicable engagement agreement.\n',encoding='utf-8')
(OUT/'START HERE.md').write_text('''# Axiomotl Advisory Branding Kit v2

Start with **01 Guidelines/Axiomotl Brand Guidelines.pdf** for the visual reference, or open **01 Guidelines/Brand Guidelines.html** for the offline web guide.

This kit follows the current purple-and-teal website identity, Lato typography and supplied stacked axolotl logo. It replaces the earlier working palette for new brand material. It does not modify the live website.

## Included
- 16-page PDF guide, editable Markdown and offline HTML
- Current transparent logo master and usage limits
- Colour values, contrast checks, importable palette and licensed Lato fonts
- Five current workflow icons, including the supplied teal Design icon
- Four editable Word documents and an eight-slide editable PowerPoint
- Two HTML email signatures, plain-text signature and email copy
- Four editable SVG social templates with PNG exports
- Two-sided business card PDF with trim and bleed boxes
- Messaging library, source notes and CSS/JSON developer tokens

## Before using
Install Lato. Replace bracketed placeholders. Confirm names, claims and commercial terms. Test email rendering. Proof print artwork with your printer.

The supplied logo is a 500px raster master. True vector artwork and reversed or small-format variants still need the original designer's source; this kit does not invent replacement logos. New production rules and template layouts are proposals for owner review.
''',encoding='utf-8')
manifest={str(p.relative_to(OUT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in OUT.rglob('*') if p.is_file() and p.name!='MANIFEST.json'}
(OUT/'MANIFEST.json').write_text(json.dumps(manifest,indent=2))
print(json.dumps({'folder':str(OUT),'files':len(manifest),'guidePages':16,'wordTemplates':len(docs),'slides':len(prs.slides)}))
